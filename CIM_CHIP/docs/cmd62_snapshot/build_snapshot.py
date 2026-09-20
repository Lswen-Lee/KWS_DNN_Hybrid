"""Reconstruct the pre-command-62 TB snapshot, not a simulator memory dump."""
from pathlib import Path
import csv
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import ListedColormap
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

out=Path(__file__).resolve().parent
mem=[[0]*32 for _ in range(128)]
tag=[[0]*32 for _ in range(128)]
def write(row,pair,word,kind):
    mem[row][pair*2]=word>>4
    mem[row][pair*2+1]=word&15
    tag[row][pair*2]=tag[row][pair*2+1]=kind
write(45,9,0xA5,1); write(46,10,0xFF,1); write(47,11,0x00,1)
for i in range(20):
    row,pair=divmod(40*16+14+i,16)
    write(row,pair,(((row&15)^pair)<<4)|(((row>>3)+pair)&15),2)
for r in range(32): write(r,2,(7*r+35)&255,3)
assert sum(x!=0 for row in tag for x in row)==110
detail=[]
for r in range(32):
    w=16*mem[r][4]+mem[r][5]; a=(5*r+3)&15
    detail.append([r,30+r,f'{0xA0|a:02X}',a,f'{w:02X}',w,mem[r][4],mem[r][5],a*mem[r][4],a*mem[r][5],a*w])
totals=[sum(d[i] for d in detail) for i in (8,9,10)]
assert totals==[2124,1744,35728]
planes=[]; e=o=0
for bit in range(7,-1,-1):
    pe=sum(((d[3]>>bit)&1)*d[6] for d in detail)
    po=sum(((d[3]>>bit)&1)*d[7] for d in detail)
    e=2*e+pe; o=2*o+po
    planes.append([bit,pe,po,e,o])
assert [e,o]==totals[:2]
header=['Row']+[f'Bank{b}' for b in range(32)]
rows=[[r]+[f'{v:X}' for v in mem[r]] for r in range(128)]
with (out/'before_cmd62_128x32.csv').open('w',newline='') as f:
    csv.writer(f).writerows([header]+rows)
colors=['FFFFFF','FFE6BF','DDEBFF','CDEFE4']
wb=Workbook(); ws=wb.active; ws.title='128x32 HEX'
ws.append(['每格=4-bit Bank的hex值；灰白=Reset；橘=命令2/4/5；藍=Burst9；綠=命令30~61'])
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=33)
ws.append(header)
for r,row in enumerate(rows):
    ws.append(row)
    for b in range(32):
        c=ws.cell(r+3,b+2); c.fill=PatternFill('solid',fgColor=colors[tag[r][b]])
        c.alignment=Alignment(horizontal='center')
for c in ws[2]: c.font=Font(bold=True)
ws.freeze_panes='B3'; ws.auto_filter.ref='A2:AG130'
for c in range(1,34): ws.column_dimensions[get_column_letter(c)].width=8
dh=['Row','Write命令','輸入byte(hex)','有效A4','W8(hex)','W8(dec)','Even W4','Odd W4','A*Even','A*Odd','A*W8']
ds=wb.create_sheet('CMD62 32-row MAC'); ds.append(dh)
for row in detail: ds.append(row)
ds.append(['TOTAL']+['']*7+totals); ds.freeze_panes='A2'
for c in range(1,12): ds.column_dimensions[get_column_letter(c)].width=19
ps=wb.create_sheet('Bit planes'); ps.append(['plane','Even partial','Odd partial','Even ACC','Odd ACC'])
for row in planes: ps.append(row)
for c in range(1,6): ps.column_dimensions[get_column_letter(c)].width=20
wb.save(out/'before_cmd62_and_CIM.xlsx')

fig,ax=plt.subplots(figsize=(18,46),dpi=150)
fig.subplots_adjust(left=.05,right=.99,top=.965,bottom=.01)
ax.imshow(np.array(tag),cmap=ListedColormap(['#ffffff','#ffe6bf','#ddebff','#cdefe4']),vmin=0,vmax=3,aspect='auto')
ax.set_xticks(range(32),[f'B{b}' for b in range(32)],fontsize=8)
ax.set_yticks(range(128),range(128),fontsize=8)
ax.xaxis.tick_top(); ax.set_ylabel('Row',fontsize=13)
ax.set_xticks(np.arange(-.5,32,1),minor=True); ax.set_yticks(np.arange(-.5,128,1),minor=True)
ax.grid(which='minor',color='#cbd4dd',lw=.35); ax.tick_params(which='minor',length=0)
for r in range(128):
    for b in range(32):
        ax.text(b,r,f'{mem[r][b]:X}',ha='center',va='center',fontsize=8,color='#12324b' if tag[r][b] else '#8e9aa6')
for r in (31.5,63.5,95.5): ax.axhline(r,color='#18324b',lw=2)
fig.suptitle('Before command 62 | 128 rows x 32 four-bit banks\nHEX values | Orange: writes 2/4/5 | Blue: burst 9 | Green: writes 30-61 (CIM target)\nBank 2p = W[7:4], Bank 2p+1 = W[3:0] | White zero: reset value',fontsize=15,y=.994)
fig.savefig(out/'before_cmd62_128x32.png'); plt.close(fig)

def table(head,data):
    return '| '+' | '.join(map(str,head))+' |\n| '+' | '.join(['---']*len(head))+' |\n'+''.join('| '+' | '.join(map(str,r))+' |\n' for r in data)
md='''# 命令62前的權重表與A4W8解說

此表根據目前 tb_dcim_pin_master.sv 的初始化與命令1～61重建，**不是實際模擬器匯出的記憶體快照**。Read不修改權重。

## 1. 如何閱讀128×32表格

- 直向Row0～127；橫向Bank0～31。每格是一個4-bit十六進位值，不是單一1-bit Cell。
- 32 Banks × 4 bits = 128個實體bit欄，所以仍是128×128-bit陣列。
- Pair p的高4 bits在Bank 2p，低4 bits在Bank 2p+1。
- Reset後全為0；命令62前共寫55個不同W8位置，即110個4-bit格子。寫入00也算已寫，PNG／Excel用顏色區分。
- 命令64以後的Group3權重尚未寫入；Row96～127此時全為0。

[完整PNG表格](./before_cmd62_128x32.png) ｜ [Excel：權重表、32列乘加、bit-plane](./before_cmd62_and_CIM.xlsx) ｜ [CSV](./before_cmd62_128x32.csv)

| 已執行寫入命令 | 位置 | 內容／圖表色彩 |
|---|---|---|
| 2 | Row45、Bank18/19 | A/5（橘） |
| 4 | Row46、Bank20/21 | F/F（橘） |
| 5 | Row47、Bank22/23 | 0/0（橘，確實寫入0） |
| 9 | Row40 Bank28～31；Row41所有Banks；Row42 Bank0～3 | 20個W8固定Burst圖樣（藍） |
| 30～61 | Row0～31、Bank4/5 | W[r]=(7r+35) mod 256（綠） |

以下每格均為hex；未寫位置顯示Reset值0：

'''+table(header,rows)+'''
## 2. 命令62選擇哪些權重？

命令62：operation=2、activation_is_a8=0、row_addr=0、bank_addr=2。
row_addr[6:5]=0選Group0的32列；Pair2選Bank4（Even）與Bank5（Odd）。其他已寫區域不參與本次計算。

```text
32個有效A4 → Group0：Row0～31
             ├─ Bank4：高4-bit權重 → Even Adder Tree → Even ACC
             └─ Bank5：低4-bit權重 → Odd Adder Tree  → Odd ACC
```

這一筆命令計算**一個32項dot product**，不是32個輸出，也不是128列全部相加。
Even與Odd是同一份W8結果的兩個nibble貢獻，並非兩個不同filter。

## 3. 測資與32列乘加明細

令r=0～31：W[r]=(7r+35) mod 256，A[r]=(5r+3) mod 16。
輸入byte=0xA0 OR A[r]，刻意放入非零高4 bits；A4模式只使用低4 bits。

'''+table(dh,detail)+table(['合計Even','合計Odd','合計W8'],[totals])+'''
範例：Row0的W=0x23=35、A=3；Even貢獻3×2=6、Odd貢獻3×3=9，重組16×6+9=105=3×35。
Row1的W=0x2A=42、A=8；Even=16、Odd=80，重組16×16+80=336。

## 4. 實際硬體如何得到结果？

硬體每次只送一個activation bit-plane。每個plane用8 clocks載入32 bits，再一個Sample上升緣接受兩路9-bit partial sum：

```text
P_even[b] = Σ ((A[r] >> b) & 1) × EvenWeight[r]
P_odd[b]  = Σ ((A[r] >> b) & 1) × OddWeight[r]
ACC_next = (ACC << 1) + P[b]
```

A4仍送8個planes（7→0），前4個為0，後4個才是A[3:0]。菱形判斷沒有額外clock。

'''+table(['plane','Even partial sum','Odd partial sum','Even ACC（sample後）','Odd ACC（sample後）'],planes)+'''
最後得到：Even=2124、Odd=1744；W8結果=(2124<<4)+1744=35728。
重組前需擴位；此處為unsigned W8，不包含signed修正、zero-point、bias與requantization。

Chip以sume／sumo同時送出兩個18-bit序列，Master接收raw結果；完整W8合併由本TB計算actual_w8，用來比較expected_w8，並非Chip額外輸出腳位。

## 5. Demo建議講法

1. 「先看全陣列：橘色是基本寫入，藍色是Burst，綠色是這次CIM所需權重。」
2. 「命令62只選Group0與Pair2，因此只讀取綠色的Row0～31、Bank4/5做運算。」
3. 「每筆8-bit權重分成高4 bits與低4 bits；兩路各有一顆Adder Tree和Accumulator。」
4. 「以Row0為例，A=3、W=35，拆成3×2與3×3；最後高位貢獻乘16再加低位，即105。」
5. 「32列同時做一個bit-plane的運算，跨8個planes左移累加後，Even=2124、Odd=1744。」
6. 「Master接收兩路序列結果，TB重組得到35728，與直接ΣA×W一致。」

整筆延遲=1 Reset+8×(8 Load+1 Sample)+18 Receive=91 clocks，50 MHz下1.82 µs；不含事前權重寫入。
建議簡報展示選定的32列乘加明細；完整128列大表適合放大查看或當附錄，不宜縮成一頁小字。
命令62後的命令63會讀Row0/Pair2並檢查仍為0x23；本表是命令62之前，不包含後續A8W8權重。
'''
(out/'CMD62_MEMORY_AND_EXPLANATION.md').write_text(md)
print('PASS: 55 W8 writes / 110 bank entries; sums',totals)
print('Bit planes:',planes)
