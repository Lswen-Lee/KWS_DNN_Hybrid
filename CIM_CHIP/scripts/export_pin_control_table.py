"""Export Chinese chip-pin control tables; verify literal control values against RTL.

Requires openpyxl. This is document/static consistency verification, not simulation.
"""
from pathlib import Path
from math import ceil
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
PINS = {
    'model_rst_n': '低有效模型初始化；正常為 1。不要每筆命令清除，會影響儲存內容。',
    'reset1': '高有效；清除 accumulator。',
    'reset2': '高有效；清除序列輸出狀態。',
    'mode_sel': '0：SRAM；1：CIM。',
    'wl_en': 'SRAM row 選擇使能；目前 CIM 計算時為 0。',
    'write_en': '將 buffer 中的權重寫入選定 row／bank pair。',
    'sense_en': 'SRAM 讀取擷取控制。',
    'pre_en': '讀取 capture 時為 1；模型不實現類比預充電。',
    'clk_en': '輸入 buffer 移入使能，不是外部 clk 本身。',
    'sel_next': '切換輸入路徑，提供已載入的 buffer 資料。',
    'output_en': '接受一次 CIM partial sum，不是逐 bit 輸出使能。',
    'nv_store': '保留腳位，目前 Master 固定輸出 0。',
    'nv_recall': '保留腳位，目前 Master 固定輸出 0。',
    'pd_en': '保留腳位，目前 Master 固定輸出 0。',
}
GROUPS = {
    'SRAM寫入': [
        ('ST_WRITE_LOAD_HIGH', '1', 'weight[7:4]', '輸入高 nibble', {'clk_en'}),
        ('ST_WRITE_LOAD_LOW', '1', 'weight[3:0]', '輸入低 nibble', {'clk_en'}),
        ('ST_WRITE_CELL', '1', '0000', '上升緣將 buffer 的 8 bits 寫入 cell pair', {'wl_en', 'write_en', 'sel_next'}),
    ],
    'SRAM讀取': [
        ('ST_READ_CAPTURE', '1', '0000', '上升緣載入讀取 serializer；該緣後 sumo 顯示 bit7', {'wl_en', 'sense_en', 'pre_en', 'sel_next'}),
        ('ST_READ_SHIFT', '8', '0000', '每個上升緣接收 sumo：bit7→bit0；同緣 chip 移位', set()),
    ],
    'CIM計算': [
        ('ST_CIM_RESET', '1', '0000', '只清 ACC 與輸出狀態，不清權重 cell', {'mode_sel', 'reset1', 'reset2'}),
        ('ST_CIM_LOAD', '每 plane 8，共 64', '四個 activation 的同一 plane', 'stage=7→0，plane=7→0；詳見資料時序表', {'mode_sel', 'clk_en'}),
        ('ST_CIM_SAMPLE', '每 plane 1，共 8', '0000', '接受 partial sum；第八次 sample 緣後顯示結果 bit17', {'mode_sel', 'sel_next', 'output_en'}),
        ('ST_RESULT_CAPTURE', '18', '0000', '每個上升緣同時接收 sume／sumo：bit17→bit0', {'mode_sel', 'sel_next'}),
    ],
    'Burst寫入': [
        ('ST_BURST_WAIT_DATA', '依 valid 等待', '0000', 'burst_weight_ready=1；等候下一筆權重', set()),
        ('ST_BURST_LOAD_HIGH', '每筆 1', 'weight[7:4]', '輸入高 nibble', {'clk_en'}),
        ('ST_BURST_LOAD_LOW', '每筆 1', 'weight[3:0]', '輸入低 nibble', {'clk_en'}),
        ('ST_BURST_WRITE_CELL', '每筆 1', '0000', '寫入本筆；若還有下一筆可同緣預取，資料不足則回 WAIT', {'wl_en', 'write_en', 'sel_next'}),
    ],
}


def controls(high):
    return {pin: int(pin == 'model_rst_n' or pin in high) for pin in PINS}


def verify_rtl():
    text = (ROOT / 'fpga/dcim_pin_master.sv').read_text()
    moore = text.split('// Moore-style pin control', 1)[1]
    defaults, cases = moore.split('case (state)', 1)
    actual = {}
    for pin in PINS:
        value = re.search(r'chip_' + pin + r'\s*=\s*([^;]+);', defaults)[1].strip()
        if value not in {'~rst', 'rst', "1'b0", "1'b1"}:
            raise ValueError(f'Unexpected default for {pin}: {value}')
        actual[pin] = int(value in {'~rst', "1'b1"})  # normal operation rst=0
    for states in GROUPS.values():
        for state, _, _, _, high in states:
            block = re.search(state + r': begin(.*?)(?=\n            (?:ST_|default:))', cases, re.S)
            values = actual.copy()
            if block:
                for pin, bit in re.findall(r"chip_(\w+)\s*=\s*1'b([01]);", block[1]):
                    if pin in values:
                        values[pin] = int(bit)
            assert values == controls(high), f'RTL control mismatch: {state}'
    print('PASS: 13 state control vectors match RTL defaults/overrides (rst=0).')


def export():
    verify_rtl()
    wb = Workbook()
    wb.remove(wb.active)
    tables = {}

    def table(name, header, rows):
        tables[name] = [header, *rows]
        ws = wb.create_sheet(name)
        for row in tables[name]:
            ws.append(row)
        ws.freeze_panes = 'C2' if name in GROUPS else 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(name='Microsoft JhengHei', color='FFFFFF', bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Microsoft JhengHei', size=11)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if cell.value == 1:
                    cell.fill = PatternFill('solid', fgColor='D9EAD3')
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='EDF3F8')
            ws.row_dimensions[row[0].row].height = 46
        ws.row_dimensions[1].height = 44
        for index in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(index)].width = 29 if index == 1 else 24
        ws.column_dimensions[get_column_letter(ws.max_column)].width = 55
        for row in ws.iter_rows(min_row=2):
            lines = max(ceil(sum(2 if ord(c) > 127 else 1 for c in str(cell.value or '')) /
                             max(1, ws.column_dimensions[cell.column_letter].width - 2)) for cell in row)
            ws.row_dimensions[row[0].row].height = max(34, 17 * lines + 10)
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.print_title_rows = '1:1'
        ws.print_options.horizontalCentered = True
        ws.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4, header=.15, footer=.15)
        ws.print_area = ws.dimensions

    table('使用說明', ['項目', '內容'], [
        ['依據', 'dcim_pin_master.sv 與 dcim_chip_pin_top.sv；目前數位 RTL，不是實體晶片電性規格。'],
        ['方向', '以 chip 為基準：輸入由 Master 驅動；sume／sumo 由 chip 回傳。clk 是共用時脈。'],
        ['前提', '各操作表皆為 rst=0、model_rst_n=1；clk 持續運行，表列動作在上升緣發生。綠底 1 表示高電位。'],
        ['初始化', 'Master rst=1 且狀態已重設：model_rst_n=0，reset1=reset2=1；其他控制=0，din=0，地址=0。初始化會清 cell。'],
        ['Idle／Done', 'rst=0 時 model_rst_n=1；其他單 bit 控制全部=0，din=0；地址保留。輸出可能保留舊值，不採樣。'],
        ['命令接受', '僅 ST_IDLE 且 start=1 的上升緣接受命令並保存參數；ST_DONE 時 busy=0，但不接受新命令。'],
        ['cycle 定義', '從命令被接受的上升緣，到進入 ST_DONE：寫3、讀9、CIM91；不包含下一筆命令接受間隔。'],
        ['命令間隔', 'ST_DONE 下一緣回 IDLE，再下一緣才可接受 start；請勿只因 busy=0 就在 DONE 發出單拍 start。'],
        ['模型限制', 'pre_en 保留但沒有類比預充電功能；nv_store／nv_recall／pd_en 固定0且未模擬NV／電源行為。'],
        ['驗證範圍', '產表腳本核對13個狀態的控制值與RTL；Excel重新載入逐格核對。非功能模擬或Vivado驗證。'],
    ])
    table('腳位定義', ['Chip 腳位', 'Master 對應', '方向（以 chip 為準）', '位寬', '功能'],
          [['clk', '共用 clk（不是 chip_clk_en）', '輸入', 1, '持續時脈，上升緣運作。']]
          + [[p, 'chip_' + p, '輸入', 1, desc] for p, desc in PINS.items()]
          + [['row_addr', 'chip_row_addr', '輸入', 7, 'SRAM：row 0..127；CIM：[6:5] 選 group 0..3。'],
             ['bank_addr', 'chip_bank_addr', '輸入', 4, 'pair 0..15；pair b 對應 even bank 2b 與 odd bank 2b+1。'],
             ['din', 'chip_din', '輸入', 4, '權重 nibble 或同一 plane 的四個 activation bits。'],
             ['sume', 'chip_sume', '輸出', 1, 'CIM even／高權重 nibble 的序列結果；SRAM不使用。'],
             ['sumo', 'chip_sumo', '輸出', 1, 'SRAM序列讀回，或CIM odd／低權重 nibble 序列結果。']])
    for name, states in GROUPS.items():
        header = ['腳位或項目', '方向'] + [s[0] for s in states] + ['備註']
        rows = [['持續 cycles', '—'] + [s[1] for s in states] + ['CIM 的 LOAD／SAMPLE 交替重複8次，不是先64次LOAD再8次SAMPLE。' if name == 'CIM計算' else '詳見使用說明的cycle起算方式。']]
        rows += [[p, 'chip 輸入'] + [controls(s[4])[p] for s in states] + [desc] for p, desc in PINS.items()]
        rows += [['clk', 'chip 輸入'] + ['持續時脈'] * len(states) + ['共用 Master／chip 時脈。'],
                 ['din[3:0]', 'chip 輸入'] + [s[2] for s in states] + ['0000 是實際驅動值，不以 X 代替。'],
                 ['row_addr[6:0]', 'chip 輸入'] + ['保存的 row_addr_q'] * len(states) + ['CIM只以[6:5]選group；讀寫使用完整7-bit。Burst每筆自動遞增。'],
                 ['bank_addr[3:0]', 'chip 輸入'] + ['保存的 bank_addr_q'] * len(states) + ['整筆讀／寫／CIM期間保持；Burst先遞增pair，15之後換row。'],
                 ['sume', 'chip 輸出'] + [('接收 even bit17→0' if s[0] == 'ST_RESULT_CAPTURE' else '不採樣') for s in states] + ['由chip產生，Master不能設定0或1。'],
                 ['sumo', 'chip 輸出'] + [('接收 bit7→0' if s[0] == 'ST_READ_SHIFT' else '接收 odd bit17→0' if s[0] == 'ST_RESULT_CAPTURE' else '不採樣') for s in states] + ['讀capture緣後顯示bit7；第8次CIM sample緣後顯示bit17。'],
                 ['階段動作', '—'] + [s[3] for s in states] + ['各操作正常完成後進入 ST_DONE。']]
        table(name, header, rows)
    table('資料時序與命令', ['操作', 'Master命令（非chip pin）', '資料／時序', '結果'], [
        ['SRAM寫', 'operation=0；row、pair、command_weight', 'LOAD_HIGH→LOAD_LOW→WRITE_CELL；3 cycles', 'done=1，result_valid=0；可另送讀命令驗證。'],
        ['SRAM讀', 'operation=1；row、pair', 'CAPTURE→SHIFT×8；9 cycles。CAPTURE緣後bit7可見，下一緣Master接收。', 'sram_read_data[7:0]；done=result_valid=1。'],
        ['CIM', 'operation=2；group取row[6:5]、pair、activation_flat、activation_is_a8', 'RESET→(LOAD×8→SAMPLE)×8→RESULT_CAPTURE×18；91 cycles', 'even_raw_result與odd_raw_result各18 bits；done=result_valid=1。'],
        ['CIM plane', 'activation_flat[(row*8)+:8]=A[row]', 'plane p=7→0；每plane依序輸入rows28..31、24..27、…、0..3', 'din[lane]=A[4*stage+lane][p]，stage=7→0，lane=0..3。'],
        ['CIM A4', 'activation_is_a8=0', '仍送8planes；p=7..4傳0，p=3..0傳activation低4bits', '與A8同為91cycles，並非縮成4次sample。'],
        ['CIM A8', 'activation_is_a8=1', '傳送每個activation的bit7..bit0', 'unsigned原始結果；W8重建為odd+(even<<4)，須先擴位避免截斷；本Master未做此重建。'],
        ['Burst寫', 'operation=3；長度1..2048且不得超出row/pair地址空間', '資料連續且第一筆與start同緣接受時3N cycles；缺資料則等待', 'valid與ready在上升緣同時1才接受；WRITE_CELL可能預取下一筆，不要混淆請求地址與chip當前寫入地址。'],
        ['錯誤', 'Burst長度0或超出2048個word的地址範圍', '命令後進入DONE，不做cell write', 'error=1、result_valid=0。'],
    ])
    output = ROOT / 'docs/DCIM_PIN_CONTROL_TABLE.xlsx'
    wb.save(output)
    check = load_workbook(output)
    for name, rows in tables.items():
        assert list(check[name].values) == [tuple(r) for r in rows], name
    md = ['# DCIM 操作控制表', '', 'Excel： [DCIM_PIN_CONTROL_TABLE.xlsx](DCIM_PIN_CONTROL_TABLE.xlsx)', '',
          '以 chip 為方向基準，正常操作 rst=0。以下與 Excel 同源生成；RTL 未修改。', '']
    for name, rows in tables.items():
        md += ['## ' + name, '', '| ' + ' | '.join(map(str, rows[0])) + ' |', '| ' + ' | '.join(['---'] * len(rows[0])) + ' |']
        md += ['| ' + ' | '.join(str(v).replace('|', '\\|').replace('\n', '<br>') for v in row) + ' |' for row in rows[1:]]
        md.append('')
    (ROOT / 'docs/DCIM_PIN_CONTROL_TABLE.md').write_text('\n'.join(md), encoding='utf-8')
    print(f'PASS: Excel round-trip, {len(tables)} sheets; saved {output} and Markdown.')


if __name__ == '__main__':
    export()
