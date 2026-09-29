"""One-page Chinese pin-control matrix: XLSX, PNG and single-page PDF."""
import argparse
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from export_pin_control_table import ROOT, GROUPS, controls, verify_rtl


def export(font_path):
    verify_rtl()
    headers = ['操作', '狀態（省略 ST_）', 'cycles', 'mode_sel', 'reset1\nreset2',
               'wl_en', 'write_en', 'sense_en', 'pre_en', 'clk_en', 'sel_next',
               'output_en', 'din[3:0]', '功能／輸出資料']
    functions = {
        'ST_WRITE_LOAD_HIGH': '載入權重高 4 bits',
        'ST_WRITE_LOAD_LOW': '載入權重低 4 bits',
        'ST_WRITE_CELL': '將 buffer 的 8-bit 權重寫入 cell pair',
        'ST_READ_CAPTURE': '載入讀取 serializer；該上升緣後 sumo 顯示 bit7',
        'ST_READ_SHIFT': '每上升緣接收 sumo：bit7→0；sume 不使用',
        'ST_CIM_RESET': '清除 ACC／序列輸出，不清權重',
        'ST_CIM_LOAD': '載入同一 plane 的32個 activation bits\n每 cycle 4 bits，stage=7→0',
        'ST_CIM_SAMPLE': '接受一次 partial sum\n第8次 sample 緣後呈現結果 bit17',
        'ST_RESULT_CAPTURE': '同時接收 sume／sumo：bit17→0\neven／odd 各18 bits',
    }
    rows = []
    for name in ['SRAM寫入', 'SRAM讀取', 'CIM計算']:
        for state, cycles, din, _, high in GROUPS[name]:
            c = controls(high)
            if state == 'ST_CIM_LOAD':
                cycles = '8/plane'
                din = '四個 A 的\n同一 bit-plane'
            if state == 'ST_CIM_SAMPLE':
                cycles = '1/plane'
            rows.append([name, state.removeprefix('ST_'), cycles, c['mode_sel'],
                         f"{c['reset1']} / {c['reset2']}", *[c[p] for p in
                         ['wl_en', 'write_en', 'sense_en', 'pre_en', 'clk_en', 'sel_next', 'output_en']],
                         din.replace('weight', 'W'), functions[state]])
    notes = [
        '共用設定：正常操作 model_rst_n=1；clk 持續運行；nv_store=nv_recall=pd_en=0。reset1/reset2 欄依序表示兩腳位設定。',
        '地址設定：SRAM 用 row_addr[6:0]＋bank_addr[3:0]；CIM 用 row_addr[6:5] 選 group、bank_addr 選 pair。命令期間地址保持不變。',
        'CIM 順序：RESET →（LOAD×8 → SAMPLE×1）重複8個 planes → RESULT_CAPTURE×18。plane=7→0；A4 的高4個 planes 傳0。',
        'CIM din[lane]=A[4×stage+lane][plane]，lane=0..3；每 plane 先傳 rows28..31，最後 rows0..3。W8 結果由外部以 odd+(even<<4) 重建。',
        '完成時間：寫3／讀9／CIM91 cycles（命令接受緣至DONE，不含下一命令間隔）。讀capture緣後顯示bit7，下一緣才由Master接收。',
        'Idle／Done：model_rst_n=1，其餘表列單bit控制=0、din=0000，地址保留。初始化 model_rst_n=0 會清權重，勿每筆操作重設。',
        '依據：dcim_pin_master.sv／dcim_chip_pin_top.sv。pre_en 為模型保留腳位，沒有類比預充電功能；sume／sumo 由chip產生，不由Master指定0/1。',
    ]
    title = 'DCIM 操作控制總表｜SRAM 寫入・SRAM 讀出・CIM 計算'
    target = ROOT / 'docs/DCIM_PIN_CONTROL_ONEPAGE'
    wb = Workbook(); ws = wb.active; ws.title = '一頁控制總表'
    ws.merge_cells('A1:N1'); ws['A1'] = title
    ws['A1'].font = Font(name='Microsoft JhengHei', size=20, bold=True, color='17365D')
    ws.row_dimensions[1].height = 35
    ws.append(headers)
    for row in rows: ws.append(row)
    for note in notes:
        n = ws.max_row + 1; ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=14)
        ws.cell(n, 1, note); ws.row_dimensions[n].height = 30
    widths = [13, 27, 12, 12, 12, 10, 12, 12, 10, 10, 12, 13, 20, 55]
    for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Microsoft JhengHei', size=11, bold=cell.row == 2)
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center' if cell.row <= 11 and cell.column < 14 else 'left')
            if cell.row == 2:
                cell.fill = PatternFill('solid', fgColor='17365D'); cell.font = Font(name='Microsoft JhengHei', color='FFFFFF', bold=True)
            elif cell.row <= 11:
                cell.fill = PatternFill('solid', fgColor='D9EAD3' if cell.value in [1, '1 / 1'] else ('EAF2F8' if cell.row < 6 else 'FFF3DC' if cell.row < 8 else 'EEF4EA'))
    for r in range(2, 12): ws.row_dimensions[r].height = 45 if r == 2 else 50
    ws.freeze_panes = 'C3'; ws.sheet_view.zoomScale = 65
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=.2, right=.2, top=.3, bottom=.3, header=.1, footer=.1)
    ws.print_area = f'A1:N{ws.max_row}'
    wb.save(target.with_suffix('.xlsx'))
    check = load_workbook(target.with_suffix('.xlsx'))
    assert list(check.active.iter_rows(min_row=3, max_row=11, values_only=True)) == [tuple(r) for r in rows]
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt, font_manager
    font_manager.fontManager.addfont(str(font_path))
    plt.rcParams['font.family'] = font_manager.FontProperties(fname=font_path).get_name()
    fig = plt.figure(figsize=(20, 11.25)); ax = fig.add_axes([.02, .31, .96, .56]); ax.axis('off')
    fig.text(.025, .94, title, fontsize=24, weight='bold', color='#17365D')
    fig.text(.025, .897, '每列為一個操作階段；數值為實際控制電位。綠底代表 1。所有動作以 clk 上升緣為基準。', fontsize=13)
    sizes = [w / sum(widths) for w in widths]
    table = ax.table(cellText=rows, colLabels=headers, colWidths=sizes, bbox=[0, 0, 1, 1], cellLoc='center')
    table.auto_set_font_size(False); table.set_fontsize(10)
    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor('#CBD5E1')
        if r == 0:
            cell.set_facecolor('#17365D'); cell.get_text().set_color('white'); cell.get_text().set_weight('bold')
        else:
            value = rows[r-1][c]
            cell.set_facecolor('#D9EAD3' if value in [1, '1 / 1'] else '#EAF2F8' if r <= 3 else '#FFF3DC' if r <= 5 else '#EEF4EA')
    for i, note in enumerate(notes): fig.text(.025, .265 - i * .034, note, fontsize=11)
    for ext in ['png', 'pdf']: fig.savefig(target.with_suffix('.' + ext), dpi=150)
    plt.close(fig)
    print('PASS: RTL controls verified; 9 state rows and 14 columns round-trip match; single-sheet XLSX and one-page PDF/PNG exported.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--font', type=Path, required=True, help='Path to a Traditional Chinese font')
    export(parser.parse_args().font)
